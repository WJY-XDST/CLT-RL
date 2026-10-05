"""Build the illustrated Chinese implementation guide with the bundled docx runtime."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import argparse
parser=argparse.ArgumentParser()
parser.add_argument('--tutorial-v2',action='store_true')
parser.add_argument('--tutorial-v3',action='store_true')
args=parser.parse_args()
args.tutorial_v2=args.tutorial_v2 or args.tutorial_v3

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"docs"
EVAL=ROOT/"results/policy_7738_20261003"
summary=json.loads((EVAL/"summary.json").read_text())
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.footer_distance=Cm(.65)
for name,size in [("Normal",10),("Title",20),("Heading 1",15),("Heading 2",11.5),("Caption",8.5)]:
    s=doc.styles[name];s.font.name="Noto Sans CJK SC";s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),"Noto Sans CJK SC")
    s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.line_spacing=1.1
    for borders in s._element.xpath('./w:pPr/w:pBdr'):
        borders.getparent().remove(borders)
doc.styles['Normal'].paragraph_format.widow_control=True
for name in ['Title','Heading 1','Heading 2']:doc.styles[name].paragraph_format.keep_with_next=True
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run();r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,level=1)
def sub(text):return doc.add_heading(text,level=2)
def page(title):doc.add_page_break();h(title)
def fig(path,caption,width=17.1):
    q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(3);q.paragraph_format.keep_with_next=True
    q.add_run().add_picture(str(path),width=Cm(width))
    c=p(caption,'Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER
def code(text):
    q=p('');q.paragraph_format.space_after=Pt(5);q.paragraph_format.line_spacing=1.1
    r=q.add_run(text);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.5)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
def equation(text):
    q=p('');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath');mr=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;mr.append(t);m.append(mr);q._p.append(m)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            if widths:cell.width=Cm(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');marg.append(e)
            tcpr.append(marg)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7EFF5' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));tcpr.append(shading)
            q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2);q.paragraph_format.line_spacing=1.08
            rr=q.add_run(str(text));rr.font.size=Pt(9);rr.bold=i==0
        if i==0:
            rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
        cant=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(cant)
    p('').paragraph_format.space_after=Pt(1)
    return t

p('轮腿机器人开链转闭链\n与工程接入说明','Title')
p('CLT RL 自定义模型　｜　2026 年 10 月 4 日' if args.tutorial_v2 else 'CLT RL 自定义模型　｜　2026 年 10 月 3 日', 'Caption')
p('本说明记录 mine_model 从 CAD 导出的关节树，补齐左右四个转动铰链，再接入 Isaac Lab 强化学习工程的过程。内容包括孔位定位、局部铰链坐标、五连杆映射、坐标系和控制参数调整，以及实际验证结果。')
p('当前结论：闭链结构与悬空驱动已通过验证；原 7738 策略直接用于新模型时不能站稳。机械建模完成和策略可迁移是两项不同的验证。')
fig(ROOT/'results/drive_preview/viewport.png','图 1 新模型悬空驱动截图　固定基座并关闭重力　该画面没有使用 RL 策略',14.8)
table(['项目','本次模型'],[
    ['刚体与关节','15 个刚体；14 个树关节；额外 4 个闭合 revolute 铰链'],
    ['主动关节','LB、LL、LW、RB、RL、RW，共 6 个'],
    ['树内被动关节','左右 S、L1、L2、L3，共 8 个'],
    ['正式仿真格式','USD 保留闭环；附带 URDF 只描述关节树'],
    ['当前状态','闭链驱动通过；原策略迁移失败；新模型独立从零训练']],[4.0,13.2])
p('原始 ZIP 和 USDZ 的内容保留在 inputs/；正式导出保存于 exports/ 的时间版本目录。下文路径均以 mine_model 或工程根目录为起点明确标注。')

page('1 从关节树补齐物理闭环')
p('URDF 要求机器人拓扑是一棵树。CAD 中虽然各孔在默认姿态已经对齐，单纯导入关节树仍不会约束这些孔始终重合。因此保留树内关节，在 USD 中增加两个右侧铰链和两个左侧镜像铰链。')
fig(OUT/'figures/topology.png','图 2 右腿拓扑变化　实线为原关节树　红色虚线为新增约束　左侧完全镜像',16.8)
table(['闭合关节','连接孔位','转轴'],[
    ['RL3_RS_closure','RL3 末端孔 ↔ RS 末端孔','CAD −X'],
    ['RL2_RB_closure','RL2 中间孔 ↔ RB 中间孔','CAD −X'],
    ['LL3_LS_closure','LL3 末端孔 ↔ LS 末端孔','CAD +X'],
    ['LL2_LB_closure','LL2 中间孔 ↔ LB 中间孔','CAD +X']],[4.4,9.2,3.6])
sub('关节命名与父节点')
p('将导出中的 link_002_joint 统一命名为 RB_joint。RB、RL、LB、LL 的父刚体在原始关节树中已经是 base_link；构建和导出时再次校验，并写入 angleReference=base_link 元数据。RL 与 LL 不使用另一根主动杆的相对角作为输入。')
p('默认 q=0 保留 CAD 装配姿态。这里的零位不是杆件竖直向下的姿态；绝对杆方向需要叠加几何装配偏置。被动关节由闭链求解器决定，不添加独立位置驱动。')

page('2 孔位定位与 USD 铰链构建')
p('scripts/find_closure_axes.py 读取相关 STL 零件，筛选法向沿 CAD X 的端面，提取只被一个三角形使用的边，连接成孔边界，再在 YZ 平面拟合圆心。匹配两零件的同轴孔后，轴向 X 坐标取相邻板面中点。')
table(['闭合位置','CAD X mm','CAD Y mm','CAD Z mm'],[
    ['右侧 RL3 ↔ RS','−321.449799','110.908800','−225.557425'],
    ['右侧 RL2 ↔ RB','−307.449799','101.303404','−86.576178'],
    ['左侧 LL3 ↔ LS','61.950211','110.908800','−225.557425'],
    ['左侧 LL2 ↔ LB','47.950211','101.303404','−86.576178']],[5.1,4.0,4.0,4.1])
p('孔轴匹配的最大横向误差约 0.000004 mm。这只是网格拟合的一致性，不是实物加工精度，也不等于仿真运动中的闭合误差。实际动态间隙由 PhysX 刚体姿态另行测量。')
sub('把同一铰轴转换为两个刚体的局部坐标')
p('先在世界坐标建立共同铰链坐标系，将它的 Z 轴旋转到实际孔轴，再分别变换到 body0 和 body1 的局部坐标。两侧 localPos 与 localRot 必须描述同一条世界铰轴。Gf 使用行向量矩阵约定，脚本采用下式：')
code('local_frame = world_joint_frame * body_world_transform.GetInverse()\nlocalPos = local_frame.ExtractTranslation()\nlocalRot = local_frame.ExtractRotationQuat()')
sub('约束设置')
code('UsdPhysics.RevoluteJoint\naxis = "Z"\nexcludeFromArticulation = True\ncollisionEnabled = False\nbody0 / body1 = 两侧真实刚体路径\nlocalPos0 / localRot0 / localPos1 / localRot1 = 局部铰链帧')
p('excludeFromArticulation 让额外闭环位于树状 articulation 之外，由 PhysX 约束求解器共同求解。它不是删除物理约束。新增闭合铰链没有 Drive。正式 USD 移除导入器位置驱动，运行时按选定控制方式，为轮电机和隐式对照的主动腿电机重新添加指定 Drive。' if args.tutorial_v3 else 'excludeFromArticulation 让额外闭环位于树状 articulation 之外，由 PhysX 约束求解器共同求解。它不是删除物理约束。新增铰链没有 Drive，主动六电机也不保留导入器位置驱动，以免与工程的力矩控制竞争。')
sub('默认重合与运动重合分别检查')
p('导入后先检查默认孔位，再做主动关节扫角和实际 VMC 驱动。无渲染仿真中，USD 变换不一定每步同步，因此误差取自 PhysX 张量刚体姿态，而不是读取可能滞后的 USD Xform。')

page('3 物理连杆与等效五连杆')
p('物理模型保留所有杆件。控制层把同轴五连杆与平行四边形传动组合，映射成以轮轴为末端的等效五连杆，从而计算腿长、摆角和主动电机力矩。')
fig(OUT/'figures/five_bar_mapping.png','图 3 实际机构与轮轴等效机构　图中为原 CAD YZ 平面　单位 mm',15.2)
table(['符号或结构','含义'],[
    ['O / A / P / C','共同髋轴 / RL与RL1轴 / RL1与RL2轴 / RL2与RB闭合轴'],
    ['K / Q / D / W','RB与RS轴 / RL2与RL3轴 / RL3与RS闭合轴 / 轮轴'],
    ['O A P C O 主环','OA=OC=94.5 mm；AP=CP=112.5 mm'],
    ['C K D Q 平行四边形','CK=QD=115.5 mm；KD=CQ=65 mm'],
    ['轮轴等效长度','两主动臂均为 210 mm；两从动臂均为 250 mm']],[5.0,12.2])
equation('W = O + (20/9)(P − O)')
p('正运动学由两圆交点确定 P，固定使用默认装配分支；根据约束方程求导得到轮轴 Jacobian，再转换成腿长和摆角的 Jacobian。检查几何不可达、圆心重合和接近奇异的情况，并触发失效处理。')
p('等效几何不意味着质量与惯量也能简单等效。PhysX 继续使用真实 15 个刚体的质量、质心和惯量；控制层只用等效机构建立输入角到虚拟腿坐标的映射。')

page('4 坐标系和角度定义')
fig(OUT/'figures/coordinate_frames.png','图 4 由 CAD 轴向转换到工程 X 前 Y 左 Z 上的 base_link',17.0)
p('CAD 实际采用 X 横向、Y 向上、Z 向前。导出工程版本时，对 base_link 的视觉、碰撞、惯性原点以及根节点子关节原点进行旋转组合；其余杆件保留原局部坐标，保证实际装配几何不变。')
code('R = [[0, 0, 1], [1, 0, 0], [0, 1, 0]]\np_new = R @ (p_CAD − base_origin_CAD) × 0.001\nbase_origin_CAD = [−136.499794, 114.528815, 0] mm')
table(['输入','定义','符号'],[
    ['q_B 和 q_L','各自相对 base_link，CAD 装配姿态为零','左腿正；右腿镜像'],
    ['绝对主动杆方向','CAD YZ 平面 +Y 零向、朝 +Z 为正','offset + sign × q'],
    ['工程虚拟腿角 theta','相对髋轴向下为零，向后为正','atan2(−X_wheel, −Z_wheel)'],
    ['CAD 工具虚拟腿角','scripts/five_bar.py 中向前为正','与工程 theta 符号相反'],
    ['轮速','左右均将正值解释为向前滚动','原始右轮角速度取反']],[4.1,9.0,4.1])
sub('原始主动角到 VMC 力矩')
equation('q = [q_B, q_L]ᵀ     s = [L, θ]ᵀ     J(q) = ∂s/∂q')
equation('ṡ = J(q)q̇     τ = J(q)ᵀ [F, T]ᵀ')
p('左右电机轴向符号已包含在新 Jacobian 内。工程将右腿力矩再次取反的旧串联逻辑，只对 legacy_serial 生效，避免新模型被重复镜像。策略动作仍为左右各一组摆角、腿长、轮速，最终映射到 LB、LL、LW、RB、RL、RW。')
p('本次通过了自动微分虚功检查、有限差分速度检查、零位与镜像检查，并保留原串联 VMC 回归测试。')

page('5 接入工程后修改了什么')
fig(OUT/'figures/integration_pipeline.png','图 5 从原始导出到 Isaac Lab 力矩驱动的文件和控制流程',17.0)
table(['项目','原工程','新模型初次接入' if args.tutorial_v3 else '新模型配置'],[
    ['模型加载','wl.urdf → URDF 导入','版本化 mine_closed_chain.usd'],
    ['腿角和力矩映射','串联髋角与相对膝角','两个 base_link 主动角的五连杆 J'],
    ['腿力矩上限','30 N·m','10 N·m，沿用导出限值'],
    ['轮力矩上限','5 N·m','10 N·m，沿用导出限值'],
    ['轮半径和轮距','67.5 / 341 mm','60 / 419.9 mm'],
    ['腿长动作范围','0.12 至 0.26 m','0.18 至 0.36 m'],
    ['高度指令范围','0.16 至 0.20 m','0.28 至 0.32 m'],
    ['支撑前馈 每腿','50 N','46.86 N'],
    ['轮速阻尼','0.5 N·m·s/rad','0.5，由 PhysX 隐式速度 Drive 求解'],
    ['轮关节与被动关节','两轮加四腿电机','连续轮关节；8 个树内被动关节无驱动']],[4.0,6.4,6.8])
p('质量总计 10.4969 kg。原始左右轮质量不对称，右轮 0.665 kg、左轮约 0.2789 kg，按导出值保留。默认根高度约 294.19 mm；高度到向下腿长的换算偏置为 hip_z − wheel_radius = −35.346 mm。')
sub('关键文件')
p('five_bar_vmc.py 实现批量运动学和虚功映射；mine_env_cfg.py 管理新模型、执行器、几何和控制参数；wheel_legged_vmc_flat_env.py 增加按模型分支的 FK、关节映射、接触检测和奇异位形终止。完整路径见工程 wheel_legged_gym_isaaclab 包。')
p('默认任务 WheelLeggedVMC-Flat-v0 指向新模型；WheelLeggedVMC-Legacy-v0 保留原模型。新实验名 mine_wheel_legged_vmc_flat 避免自动混用旧日志。USD 与 model.json 通过 assets/robots/mine/manifest.json 选择版本。')

page('6 机械驱动测试与验证边界')
p('固定基座、悬空、关闭重力，使用两环境不同相位测试 12 秒。腿长围绕 259.19 mm 变化 ±12 mm，摆角围绕 2.99° 变化 ±3.44°，轮速变化 ±3 rad/s。物理步长 5 ms，VMC 每步计算。')
fig(ROOT/'results/bench_drive/tracking.png','图 6 悬空 VMC 的目标与实际跟随　这是机械与控制映射验证',15.6)
table(['指标','双环境结果'],[
    ['腿长 / 摆角 / 轮速平均绝对误差','0.838 mm / 0.098° / 0.0374 rad/s'],
    ['最大闭合间隙 / 轮轴 FK 误差','0.00080 mm / 0.00222 mm'],
    ['重置 / 数值失效','0 / 0']],[7.0,10.2])
p('质量矩阵由 PhysX 根据新模型重新计算：固定基座为 14×14，自由基座为 20×20，检查结果为有限、正定。它们是树广义质量矩阵，闭链约束另行求解；本次没有把它们误当成消元后的 6×6 主动坐标矩阵。')
p('解除固定基座后的重力检查中，机构约束仍闭合，但没有平衡策略时会接地。悬空图中的直立姿态由固定基座保持，不是某个 RL 策略的站立成绩。')

page('7 原训练策略在新模型上的表现')
p('使用最新 model_7738.pt，权重与输入定义不变；新旧模型分别使用各自的 VMC 配置。9 种平地工况 × 43、44、45 三个扰动种子，每次 8 秒。根速度初始扰动幅值为 0.05；启动立即以 0.8 m/s² 斜坡下发速度，转向为 ±0.3 rad/s。')
p('原模型 27/27 完成、零重置；新模型 0/27 完成，首次失效均在 0.44 至 0.48 秒，8 秒重复尝试累计 460 次机身接地重置。没有数值失效或奇异位形终止。')
table(['新模型工况','速度 MAE m/s','高度 MAE mm','首次失效 s'],[
    [r['name'].replace('forward_0p3','前进 0.3').replace('reverse_0p3','后退 0.3').replace('stand','站立'),f"{r['all_sample_speed_mae_m_s']:.3f}",f"{r['all_sample_height_mae_mm']:.1f}",'0.44 至 0.48' if r['name']=='forward_0p3' else '0.46'] for r in [summary['mine']['conditions'][i] for i in [0,1,3]]
],[5.2,4.0,4.0,4.0])
p('表中 MAE 包含反复失败过程，误差相对实际施加的斜坡指令；不表示稳态跟踪精度。±0.6 m/s 工况最高只施加到约 +0.384 / −0.368 m/s 就重置，未完成目标速度验收。')
fig(EVAL/'startup_errors.png','图 7 新模型首次失效前的误差　叉号为终止点　三条线对应三个种子',14.6)
p('站立工况平均约 34.3% 的执行器采样处于力矩限幅，23% 左右的策略动作分量被裁剪。新模型的腿长、质量、轮速控制增益、动作范围与训练模型不同；新高度指令也超出原训练 0.16 至 0.20 m 的范围。现有证据证明直接迁移失败，尚未将原因归结为某个单一参数。')
p('完整 8 秒误差曲线、左右腿误差、力矩限幅曲线、原模型对照和逐工况统计位于 results/policy_7738_20261003/。原始数据在重置前采集，图中保留失败点并断开重置连接。')

page('8 整理后的目录与复现入口')
p('mine_model 根目录只保留总说明和统一运行入口。源数据、正式导出、有效测试和历史调试分开；移动前后的路径记录在 docs/relocation_map.json，原始 ZIP 与正式 USD 哈希校验一致。')
code('mine_model/\n  README.md              总入口和当前状态\n  run.sh                 统一运行命令\n  inputs/                原始 ZIP 与 USDZ\n  source/                原始解压文件\n  prepared/              命名规范化的导入输入\n  config/                闭合铰链与孔位拟合结果\n  scripts/               构建 检查 导出 计算和分析工具\n  exports/               正式模型 按时间版本保存\n  results/\n    bench_drive/         悬空驱动通过记录\n    gravity_check/       重力检查 不是平衡通过\n    drive_preview/       悬空测试截图\n    policy_7738_20261003/ 原策略在新旧模型上的评估\n  docs/                  本说明 配图和详细实现记录\n  archive/               失败调试及被替代的评估\n  build/                 可再生场景 缓存和日志')
sub('在工程根目录运行')
code('./mine_model/run.sh inspect\n./mine_model/run.sh bench --num_envs 1 --loop\n./mine_model/run.sh test\n./mine_model/run.sh build --headless --steps 240\n./mine_model/run.sh export\n./mine_model/run.sh evaluate --headless --robot mine\n./mine_model/run.sh evaluate --headless --robot legacy')
p('inspect 用于人工查看和调主动 joint 角度；bench 是固定基座驱动，不能当作站立策略。evaluate 加载 7738 策略在自由基座上评估。export 创建新的时间版本并更新工程资产 manifest，不覆盖原始 CAD 文件。')
sub('结果与源文件索引')
p('正式模型：exports/export_20261003_181515/。完整闭链文件为 mine_closed_chain.usd；附带的 mine_robot.urdf 需要 meshes/，且不能单独保存四个额外闭环。')
p('机械证据：results/bench_drive/report.json、trace.csv、tree_mass_matrix.npy。策略证据：results/policy_7738_20261003/summary.json、summary.csv、error_curves.pdf，以及 mine/、legacy/ 中的 manifest、resolved_env 和 trace。')
page('9 新模型参数核对与独立训练')
p('原策略评估过程中没有更新旧权重。评估完成后，按要求单独启动新模型从零训练，旧 checkpoint 与旧实验目录保留。此前所见能保持直立的悬空演示由固定基座实现，与新策略训练成绩无关。')
table(['参数','实际使用的新结构值'],[
    ['质量与惯量','USD 内 15 个刚体的 CAD 质量、质心和惯量；总质量 10.4969 kg'],
    ['几何','94.5 / 112.5 mm 主环，经 20/9 放大到轮轴'],
    ['轮半径与轮距','60 / 419.9 mm'],
    ['坐标系和电机参考','X 前 Y 左 Z 上；4 个髋电机均相对 base_link'],
    ['控制','mine_five_bar 分支；左右符号、FK、J 与力矩限幅均使用新模型'],
    ['初始增益','接入时 Kpθ=50、Kdθ=3、KpL=900、KdL=90；最终值见策略接口契约']],[5.1,12.1])
p('旧配置中的 l1=0.15、l2=0.25、offset=0.054 仍保留，用于 Legacy 任务的串联分支。新任务在 mine_five_bar 分支直接返回，读取 model.json 的几何；这些旧字段不会参与新模型腿运动学或力矩计算。')
sub('训练设置与存储位置')
code('任务  WheelLeggedVMC-Flat-v0\n模型  新建闭链 USD\n起点  随机初始化  不加载旧 checkpoint\n环境  256 个\n预算  1000 次 PPO 迭代  每迭代 48 步\n种子  43\n保存  每 100 次迭代保存 checkpoint')
code('IsaacLab/logs/rsl_rl/\n  wheel_legged_vmc_flat/       原模型训练记录\n  mine_wheel_legged_vmc_flat/  新模型独立训练记录')
p('32 环境、2 次迭代的短程检查已完成，成功保存 model_0.pt 与 model_1.pt。正式训练使用独立新运行，从零开始；训练启动不代表学会平衡，应根据新 checkpoint 在自由基座上的站立和运动曲线另行验收。')
sub('运行与检查')
code('./train_mine.sh --headless\ntail -f mine_model/build/logs/training_mine_formal.log')
p('历史试训已经到第 999 轮并评估，结果未通过。随后优化到第 1800 轮；发现蹲塌奖励漏洞后保留 checkpoint，修正判定和奖励，再继续优化。训练记录与评估证据保留在各自时间目录，不能将“进程正常”视为站立成功。')

if args.tutorial_v2:
    page('10 站立失败的定位和修正')
    p('先检查角度与关节，再检查自由基座的负载控制，最后检查奖励和终止。CAD 电机 q=0 时，虚拟腿长为 259.1946 mm，虚拟腿角约 +0.05222 rad，即 2.99°；它不等于虚拟腿竖直零角。两种零点已经通过装配偏置显式换算，不能在策略回放时随意再减一个角度。')
    fig(OUT/'figures/balance_diagnosis_reward.png','图 8 真实蹲塌轨迹与高度惩罚修正　30 N·m 只用于诊断',15.5)
    p('第 1600 轮策略的 5 秒零扰动站立诊断中，目标腿长约 0.29 m，实际约 0.14 m，机身约 0.14 m。给 B 或 L 电机正力矩会产生正确方向的关节运动，径向 47 N 会使悬空腿伸长；带重力的 FK 轮轴位置误差约毫米级，没有发现大幅零点或镜像反向。提高至 30 N·m 仍塌陷，不能仅用“力矩上限偏小”解释。')
    p('旧高度平方误差权重为 −1000，但每秒单项惩罚被限制为 −1，超过 31.6 mm 就没有进一步区别。机身接触依赖连续 0.15 秒超过 10 N，间歇接触可以让蹲塌回合持续。现在扩大高度惩罚范围，加入站立绝对腿角项，并以机身低于 0.22 m 持续 0.15 秒作为独立失败条件。')
    p('这些结果确认了奖励与终止漏洞，尚不能证明它是唯一原因。悬空驱动、力矩脉冲和短传统反馈对照都不是自由基座 RL 站立验收。新版训练必须通过实际平地回放，原来的蹲塌 checkpoint 不作为可运行成果。')

    page('11 自动优化和百分之二十验收')
    p('使用 256 个环境；新隐式策略试训 1000 轮，后续每 200 轮回放。先站立后运动，按实测超标项调整；固定偏航测试不加航向纠偏，显式 VMC 失败记录保留。' if args.tutorial_v3 else '正式任务仍使用 256 个并行环境。由新模型第 1800 轮 checkpoint 继续训练，保留旧记录；每 1000 轮评估一次。先满足站立精度，再加入前进、后退、转向和高度变化，依据最差工况调整奖励或延长站立课程。')
    p('第 999 轮已 27/27 完成 12 秒、零重置，运动和轮速精度仍未通过。站立轮速 MAE 约 0.51 rad/s，因此增加静止轮速跟踪惩罚。站稳后保留优化器，运动阶段单独增加轮动作探索；实际倒地才恢复全部探索，最佳模型按全部指标排序。' if args.tutorial_v3 else '首轮到 2799 轮仍为 0/27 完成。摆角增益短测支持建立 Kpθ=10、Kdθ=0.5 的训练对照；探索 std 过早从 0.3 降到约 0.04，因此新模型熵系数由 0.001 调至 0.01。这些调整尚未证明站立通过，最终以实际回放验收。')
    table(['指标','验收要求'],[
        ['自由基座生存','9 工况 × 3 种子；首回合完整、无重置、无数值失效'],
        ['速度和偏航 MAE','非零目标 20%；接近零使用 0.05 m/s 或 rad/s 绝对容差'],
        ['高度和虚拟腿长 MAE','各自目标的 20%；稳态最低高度不低于指令的 80%'],
        ['左右腿摆角 MAE','各腿相对目标 ≤0.04 rad，取 0.2 rad 设计范围的 20%'],
        ['轮速 MAE','20% 参考轮速；接近零绝对容差 0.3 rad/s'],
        ['机身倾角','俯仰与横滚绝对角的 p95 ≤5°'],
        ['独立复验','同一权重用 83 84 85 种子，再每工况运行 30 秒']],[5.1,12.1])
    p('稳态误差取最后 3 秒，失败回合不产生稳态成绩。任何工况或种子失败，整个 checkpoint 都不能通过。零目标没有可定义的百分比误差，采用上述明确的绝对容差。')
    sub('通过后的交付和关机')
    p('通过复测后导出 TorchScript、ONNX 与完整 sim2sim_contract.json，记录相同 checkpoint 的 SHA256，并保存原始曲线和验收报告。自动生成训练更改说明，连同本教程的 DOCX 与 PDF 打包，校验压缩包和文件哈希，刷盘后按用户要求关机。测试失败或文档校验失败不会触发关机。')
    state_directory='auto_training_implicit_20pct_20261004' if args.tutorial_v3 else 'auto_training_20pct_20261004'
    code('状态目录  mine_model/results/\n          '+state_directory+'/\nstatus.json                  当前阶段和训练子进程\nround_*/                     训练配置 日志 逐轮回放\naccepted_policy/             验收通过的策略和接口契约\nfinal_delivery/              最终模型 数据和说明\ncompletion.json              完成与关机执行记录')
    p('本教程描述建模与验收流程，不代表当前训练已通过。新模型在 MuJoCo 的 sim2sim 回放属于后续独立验证；Isaac Lab 验收不能替代它。')

if args.tutorial_v3:
    page('12 闭链负载控制的执行方式对照')
    p('增益、求解迭代数、接触求解顺序、速度差分、较小步长与释放轮电机等诊断都未使显式 VMC 承重通过。另用同一 CAD q=0 目标比较显式关节 PD 与 PhysX 隐式 Drive：10 Nm 限幅下，前者约 0.30 秒失败，后者在竖直导轨中持续 3 秒，机身约 0.290 m。这定位了执行方式差异，但尚未确定 PhysX 内部的唯一原因。')
    fig(OUT/'figures/implicit_actuator_diagnosis.png','图 9 同目标同限幅的实测对照与候选控制流程　导轨限制机身倾斜和水平移动',16.7)
    sub('五连杆逆解与电机目标')
    p('先按相同动作映射得到腿长和摆角目标。将轮轴目标除以 20/9 得到主环点 P；用余弦定理分别求两根 94.5 mm 主动杆角度，选择 CAD 装配对应的肘部分支，减去装配偏置并应用左右轴符号，得到相对 base_link 的 B、L 电机目标。')
    equation('cos α = (u² + r² − d²) / (2ur)')
    p('新增 five_bar_inverse，完整腿长 0.18 至 0.36 m、摆角 ±0.2 rad 网格的正逆回代通过。候选使用 force 类型关节 Drive，初始关节 Kp=300 Nm/rad、Kd=3 Nm·s/rad，正式限幅仍为 10 Nm。8 组隐式增益在导轨中承重 5 秒均未失败；自由基座短反馈对照仍失衡，不能作为 RL 验收成果。')
    p('执行契约由 explicit_vmc 的 Jᵀ[F,T] 显式力矩，改为 implicit_joint_reference 的逆运动学关节目标与隐式 PD。FK、左右 joint 对应、坐标基准和 27 维观测 / 6 维动作保持一致。MuJoCo 端必须同步实现新关节驱动及相同限幅，不能只复制策略网络而保留旧 VMC 力矩执行方式。')

    page('13 运动奖励裁剪的检查与修正')
    p('第 1399 轮通过站立，加入运动后出现反向直行。第 2199 轮回放 24/27 完成，前进 0.6 m/s 的速度 MAE 最差约 1.091 m/s，高机身失败。轮轴与 Drive 符号经核验一致，不能通过翻转轮速掩盖策略反向响应。')
    fig(OUT/'figures/reward_clipping.png','图 10 相同奖励公式的裁剪对照　解析曲线，不是训练后的跟踪成绩',14.5)
    sub('为什么只加权重还不够')
    p('原速度平方惩罚系数为 −30，却共用每秒 −1 的裁剪。误差超过 √(1/30)=0.182574 m/s 后，停止和反向运行均进入同一平台。继续加大系数会更早进入平台。2000 轮候选首回合数据中，速度项 43.06%、静止轮速项 93.10% 的采样达到旧上限。')
    table(['新增配置','2400 轮训练值 reward/s','覆盖的误差'],[
        ['lin_vel_penalty_clip','43.2','速度误差 1.2 m/s'],
        ['yaw_rate_penalty_clip','2.8125','偏航误差 0.6 rad/s'],
        ['standing_wheel_penalty_clip','12.20703125','左右轮各 2 rad/s']],[6.3,5.3,5.6])
    p('新增项默认上限 1，自动优化按权重扩大，保留 Legacy 行为。静止轮速项仅在速度及偏航静止时比较实际轮速与策略参考。修改训练奖励，观测、五连杆逆解和执行器契约不变。')
    sub('检查过程与复现入口')
    p('先从 trace.csv 计算旧上限饱和比例，再核对每轮 overrides.json 的系数与裁剪；训练加入实际命令、实际速度及 Reward_Clipping 日志。站立存活时比较全部工况的超标比例，避免静止项遮住更大的运动误差。high_height 日志阈值也由旧 0.19 m 改为新范围的 0.31 m。')
    code('results/auto_training_implicit_20pct_20261004/\n  reward_saturation_audit_20261004.json\n  round_*/replay/  验收报告及全部误差曲线')
    p('裁剪修正仍需完整回放证明有效。最终更改说明自动记录通过验收的 checkpoint、实际参数、三种子最差误差和复现命令；没有通过全部验收与文件校验时不关机。')

    h('14 策略初始化与高度参考的实测对照').paragraph_format.page_break_before=True
    p('扩大奖励后，原策略仍持续反向直行。机械轮轴和轮速符号经复核正确，因此另建可追溯的策略初始化副本。以下对照固定 2199 轮源权重，使用完整指令、9 工况、3 个种子、每工况 12 秒。所有修改均在训练前完成，回放期间不改权重或观测。')
    fig(OUT/'figures/policy_initialization_comparison.png','图 11 初始化与控制对照的实测误差　最终候选仍未达到全部验收指标',14.5)
    sub('速度方向与腿摆角初始化')
    p('把 actor 和 critic 首层的速度命令输入列 6 取负，并同步变换 Adam 一阶矩；二阶矩、探索标准差和其它权重保留。这只生成续训起点，运行时输入正常的 27 维观测。随后将 actor 输出行 0 和 3 改为 25% 原输出与 75% CAD 名义腿角的混合；名义归一化角为 0.2610973，只清除这两行的优化器历史。')
    sub('高度外环与轮阻尼')
    p('轮速阻尼由 0.5 提至 2.0；腿长目标以 95% 高度命令对应的名义腿长和 5% 策略腿长混合，再加入 0.25 倍高度误差修正。最终目标仍经物理限幅和五连杆逆解。这是执行器参考契约的调整，MuJoCo 必须同步；不改变 CAD 质量、关节轴或被动铰链。')
    table(['对照','存活','直行姿态 p95','0.6 m/s 前进 MAE'],[
        ['原 2199 轮策略','24 / 27','约 11 至 16 度','1.091 m/s'],
        ['仅速度列初始化','24 / 27','约 11 至 16 度','0.132 m/s'],
        ['加腿角初始化及高度参考','27 / 27','全部不超过 5 度','0.172 m/s']],[6.8,2.0,4.2,4.2])
    p('最终候选无重置，高度 MAE 约 2 至 5 mm；仍有速度、偏航及静止轮速超标，故仅作为正式续训与回退起点。保留原策略、修改清单和 checkpoint 哈希。后续优化继续完整指令验收，达标后再独立复测与交付。')
    code('results/speed_response_initialization_2199_20261004/\n  initialization.json          速度列修改来源与哈希\n  angle_0p25/initialization.json  腿摆角修改清单\n  angle_0p25/damping_2p0_height_reference/replay/')

    h('15 轮动作快速摆动与平滑奖励').paragraph_format.page_break_before=True
    p('2800 轮回放 27/27 存活，只有偏航未通过。3000 轮增加偏航权重仍未改善。原始曲线中，轮速参考约 22 至 25 Hz 反复变化，接近 50 Hz 控制的奈奎斯特频率；偏航误差同时包含慢漂移与快速振荡。')
    fig(OUT/'figures/wheel_flutter.png','图 12 左两图来自站立种子 43 的实测　右图为奖励公式对照',14.5)
    sub('先保持权重做阻尼对照')
    p('固定 2799 轮策略，仅将轮速度 Drive 阻尼由 2.0 改为 1.0 或 0.5。两组均 27/27 存活、零重置，但没有通过全部精度，轮速 MAE 增大，因此正式阻尼保持 2.0。')
    table(['配置','前进 0.6 偏航 MAE','最差轮速 MAE','全部验收'],[
        ['阻尼 2.0','0.0901 rad/s','0.7750 rad/s','未通过'],
        ['阻尼 1.0','0.0821 rad/s','1.0362 rad/s','未通过'],
        ['阻尼 0.5','0.0770 rad/s','1.7500 rad/s','未通过']],[3.4,5.0,4.5,4.3])
    sub('用训练奖励减少大幅来回变化')
    p('action_rate 对六维归一化动作计算相邻步变化平方和，包含轮动作；action_smooth 二阶项仅包含腿。原权重 −0.01 下，两轮从 −1 跳到 +1 的代价仅为每秒 −0.08；新权重 −0.5 对应每秒 −4。')
    equation('r_rate = w_rate × Σ(a_t − a_(t−1))² × dt')
    p('新增 action_rate_penalty_clip，默认 1 保留旧行为；新对照上限 12 reward/s，覆盖六维各自最大变化 2，避免增强权重后被共用裁剪压平。只改训练奖励，观测和执行器参考不变。')
    p('3200 后增强平滑代价；精度阶段再将 PPO 熵系数从 0.01 降至 0.001，保留权重、优化器和现有探索分布。效果须经完整回放，公式图不代表收敛。最终参数与独立验收见更改说明。')
    code('results/auto_training_implicit_20pct_20261004/\n  wheel_flutter_audit_20261004.json\nresults/wheel_flutter_damping_2799_20261004/\n  damping_1p0/  damping_0p5/  配置及完整回放')

    h('16 精确反馈诊断与轮速输出精度训练').paragraph_format.page_break_before=True
    p('约 5000 至 5400 轮的正式回放均 27/27 存活，速度、高度、腿长、腿角和姿态已进入验收范围，但部分偏航及静止轮速仍超标。以下方法用于排查和训练对照，不表示已得到最终合格模型。')
    fig(OUT/'figures/policy_feedback_precision.png','图 13 精确观测诊断流程与仅训练轮速输出的 PPO 对照',15.8)
    sub('精确输入记录与受控参数对照')
    p('使用可选 record_policy_inputs 保存实际 27 维输入及策略均值，不改变控制。固定 4799 权重，CPU actor 重算与 GPU 均值最大差约 8.35e-7。自动微分用于检查局部响应，不能单独证明 PhysX 闭环稳定。只从曲线差分重建观测会产生明显误差，因此不用于确认反馈增益。')
    p('测量列 1 是俯仰角速度，列 2 是偏航角速度，列 7 是偏航命令。反馈缩放和轮输出差动缩放均另存新 checkpoint，运行时仍使用正常输入。每个候选保持来源、哈希、控制参数与完整回放；已做的对照仍有超标项，不能直接宣称验收通过。')
    sub('轮速输出精度训练对照')
    p('只更新 actor 最后层轮输出行 2 和 5、轮探索参数及 critic；冻结共享特征、腿角与腿长输出行 0、1、3、4 及对应探索参数。冻结行的 Adam 历史清零，避免零梯度仍被旧动量移动；轮与 critic 的历史保留。实时训练后逐项比较 checkpoint，确认固定参数没有移动。')
    sub('自动调整需要实测依据')
    p('惩罚达到上限时明确记录继续学习，不把相同参数记为调整。若所有工况存活、轮参考相邻步变化 RMS 超过 1 rad/s，而超标项为偏航或静止轮速，可逐轮增强 action_rate，绝对权重不超过 2，并同步扩大裁剪范围。诊断统计独立于验收误差，20% 容差与独立种子复测要求不变。')

    h('17 圆轮碰撞网格的分辨率检查').paragraph_format.page_break_before=True
    p('检查 USD 实际碰撞顶点发现，左右轮胎外轮廓均为 18 边，半径为 60 mm。显示网格看起来接近圆轮，但碰撞计算使用多个平面；每个平面中点比理想圆轮低约 0.912 mm。滚动时，支撑位置及接触法向会随棱边切换。')
    fig(OUT/'figures/tire_collision_comparison.png','图 14 实际轮廓分辨率与同权重回放对照　虚线 1 为各指标验收上限',15.8)
    sub('另存圆轮碰撞 USD 叠加层')
    p('生成 mine_closed_chain_tires_96.usd，引用原 USD 并只修改轮胎碰撞：关闭原轮胎碰撞，新增 96 边圆轮网格，沿原轮轴局部 Z 放置，保留半径与约 20 mm 轴向宽度。设置 hullVertexLimit=255，提高凸包顶点上限。原 CAD 显示网格、质量、质心、惯量、joint 轴与四个闭合铰链均保留；原文件不覆盖。')
    sub('用固定权重确认影响')
    p('固定经过精度训练的 5500 权重，保持同一指令、种子及控制参数。原碰撞仍有偏航和静止轮速超标；96 边变体首轮 27/27 全部通过、无重置。该对照支持轮胎碰撞分辨率会影响精度，不能单凭棱边频率将所有高频波动归因于网格。')
    p('5500 权重在原碰撞上训练，以上对照未改权重。独立 12 秒复测也通过，但 30 秒偏航最大 0.0602 rad/s，超过 0.05；因此在 96 边资产上继续训练轮输出和 critic。最终交付须三阶段全通过。复现须加载正式 overrides.json 指定的 USD 和控制参数；叠加层及底层源 USD 哈希同时核验。')
    code('./run_python.sh mine_model/scripts/build_smooth_tire_variant.py --sides 96')
    p('变体位于工程 assets/robots/mine/export_20261003_181515；同名 JSON 记录来源、机械参数核验和文件哈希。')

doc.core_properties.title='轮腿机器人开链转闭链与工程接入说明'
doc.core_properties.subject='闭链建模 工程控制适配 策略迁移评估 目录整理'
doc.core_properties.author=''
path=OUT/('开链转闭链图文教程_v3.docx' if args.tutorial_v3 else ('开链转闭链图文教程_v2.docx' if args.tutorial_v2 else '闭链建模与工程接入说明.docx'))
doc.save(path)
print(path)
if args.tutorial_v2:
    raise SystemExit(0)

# Keep the evaluation landing page next to its plots and raw data.
lines=['# 7738 策略在新模型上的平地评估','',
    '**结论：原模型 27/27 完成，新模型 0/27 完成；新模型首次失效约 0.44–0.48 秒。**',
    '','最新 checkpoint 为 2026-10-01 21:06 训练轮次的 model_7738.pt，完整路径与 SHA256 位于 mine/manifest.json。没有训练或调整策略权重。',
    '','9 工况 × 3 个扰动种子，每工况 8 秒，物理 200 Hz，策略 50 Hz，根初速度扰动幅值 0.05。启动即下发指令，速度按 0.8 m/s² 斜坡变化。',
    '','新模型高度 0.28/0.30/0.32 m；原模型高度 0.16/0.18/0.20 m。两者使用对应的模型/VMC 配置，不是完全相同物理条件的性能排名。',
    '','## 误差曲线','',
    '- [全部误差曲线 PDF](error_curves.pdf)',
    '- [站立与前后运动](errors_straight.png)',
    '- [较快运动与左转](errors_fast_and_turn.png)',
    '- [右转与高低姿态](errors_turn_and_height.png)',
    '- [首次失效前放大曲线](startup_errors.png)',
    '- [左右腿误差与力矩限幅](leg_and_saturation_errors.png)',
    '', '![首次失效前误差](startup_errors.png)','',
    '## 新模型逐工况统计','',
    '以下 MAE 包含全部失败与重试采样，速度误差相对实际施加的斜坡命令；不是稳态跟踪指标。所有工况均没有可用的首回合末 2 秒稳态窗口。',
    '','| 工况 | 重置总数 | 速度 MAE m/s | 高度 MAE mm | 完成数 |','|---|---:|---:|---:|---:|']
for r in summary['mine']['conditions']:
    lines.append(f"| {r['name']} | {r['resets']} | {r['all_sample_speed_mae_m_s']:.3f} | {r['all_sample_height_mae_mm']:.1f} | {r['survivors']}/3 |")
lines += ['','所有 460 次终止均包含 base_contact，没有数值或奇异位形失败。±0.6 m/s 的目标未到达：斜坡最大约 +0.384/−0.368 m/s 就重置。站立工况平均执行器采样力矩饱和比例约 34.3%。',
    '','## 数据口径','',
    'trace.csv 的每行均在物理步进后、自动重置前记录，episode 列区分每次重新启动。曲线红色竖线表示重置，重置之间断线，失败帧不删除。summary.json 区分全采样 MAE、首回合 MAE 与存活到末 2 秒的稳态 MAE；不存在的稳态指标为 null。',
    '','mine/ 与 legacy/ 包含完整配置和原始轨迹；summary.csv 可直接导入分析工具。archive/policy_settle_1s 保留了早期等待 1 秒再运动的试验，因新模型未能存活到运动阶段，没有用于本页运动成绩。',
    '','## 复现','',
    '从工程根目录运行 `./mine_model/run.sh evaluate --headless --robot mine` 和 `--robot legacy`，分别输出新时间目录。把两次输出分别放入同一目录的 mine/、legacy/ 后，运行：',
    '', '```bash','python3 mine_model/scripts/analyze_policy.py <该评估目录>','```','']
(EVAL/'README.md').write_text('\n'.join(lines))
